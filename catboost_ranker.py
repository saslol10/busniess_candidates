"""CatBoost: строгая схема, проверка модели, ранжирование без фиктивных баллов."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
from candidate_features import FEATURE_COLUMNS, SPECS, ALL_FEATURE_COLUMNS, ALL_SPECS

# v1 — 4 направления (старые модели), v2 — плюс учредители и руководители.
SCHEMA_VERSION = 2
SCHEMA_FEATURES = {1: list(FEATURE_COLUMNS), 2: list(ALL_FEATURE_COLUMNS)}
ROOT = Path(__file__).resolve().parent
DEFAULT_MODEL = ROOT / "models" / "catboost_model.cbm"
LABEL = "Метка_целевая (1/0)"
HEADERS = {"catboost_rank": "Место CatBoost", "catboost_score": "Оценка CatBoost (0–1)",
           "catboost_mode": "Режим CatBoost", "catboost_version": "Версия CatBoost-модели"}


def digest(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def prepare_features(frame, columns=None):
    """Проверка и приведение признаков. columns: список схемы (по умолчанию v1)."""
    columns = list(columns or FEATURE_COLUMNS)
    if not frame.columns.is_unique:
        raise ValueError("Названия столбцов не должны повторяться.")
    missing = [c for c in columns if c not in frame.columns]
    if missing:
        raise ValueError("Не хватает признаков: " + ", ".join(missing))
    x = frame.loc[:, columns].copy()
    for c in columns:
        x[c] = pd.to_numeric(x[c], errors="raise").astype(float)
        if np.isinf(x[c]).any() or (x[c].dropna() < 0).any():
            raise ValueError("Недопустимые значения признака " + c)
    for group, specs in ALL_SPECS.items():
        done, found = group + "_search_completed", group + "_candidate_found"
        if done not in columns:
            continue
        if x[done].isna().any() or not x[done].isin([0, 1]).all():
            raise ValueError(done + ": допустимы только 0 и 1.")
        absent = x[done].eq(0)
        cols = [found] + [s[0] for s in specs]
        if x.loc[absent, cols].notna().any().any():
            raise ValueError(group + ": при невыполненном поиске нужны пустые значения, не нули.")
        known = ~absent
        if x.loc[known, cols].isna().any().any() or not x.loc[known, found].isin([0, 1]).all():
            raise ValueError(group + ": незаполненные результаты выполненного поиска.")
        not_found = known & x[found].eq(0)
        for c, _, kind in specs:
            if (x.loc[not_found, c] != 0).any():
                raise ValueError(c + ": при отсутствии кандидата нужен ноль.")
            v = x[c].dropna()
            if kind == "percent" and (v > 100).any():
                raise ValueError(c + ": процент больше 100.")
            if kind == "count" and (v != np.floor(v)).any():
                raise ValueError(c + ": число объектов должно быть целым.")
    return x


def score_candidates(frame, model_path=None):
    """Возвращает (таблица, метаданные). Ошибка модели не ломает поиск."""
    out = frame.drop(columns=list(HEADERS), errors="ignore").copy()
    info = {"mode": "CatBoost не применён: модель не загружена", "applied": False,
            "warning": "", "version": "", "calibrated": False}
    out["catboost_score"] = np.nan
    out["catboost_rank"] = np.nan
    path = Path(model_path) if model_path else DEFAULT_MODEL
    if out.empty:
        info["mode"] = "CatBoost не применён: нет кандидатов"
    elif path.is_file():
        try:
            manifest_path = path.with_suffix(".json")
            if not manifest_path.is_file():
                raise ValueError("Рядом с .cbm необходим файл " + manifest_path.name)
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            columns = SCHEMA_FEATURES.get(manifest.get("schema_version"))
            if columns is None or manifest.get("features") != columns:
                raise ValueError("Модель и приложение используют разные схемы признаков.")
            if manifest.get("positive_label") != 1 or manifest.get("model_sha256") != digest(path):
                raise ValueError("Модель не соответствует манифесту (класс или контрольная сумма).")
            validation = manifest.get("validation_status")
            if validation not in ("held_out_test", "unvalidated"):
                raise ValueError("Манифест не содержит понятного статуса проверки модели.")
            from catboost import CatBoostClassifier, __version__
            model = CatBoostClassifier()
            model.load_model(str(path))
            if model.feature_names_ != columns or list(model.classes_) != [0, 1]:
                raise ValueError("Несовместимые признаки или классы внутри файла модели.")
            x = prepare_features(out, columns)
            out["catboost_score"] = model.predict_proba(x)[:, 1]
            if not np.isfinite(out["catboost_score"]).all():
                raise ValueError("Модель вернула некорректные оценки.")
            out = out.sort_values(["catboost_score", "candidate_inn"], ascending=[False, True],
                                  kind="mergesort").reset_index(drop=True)
            out["catboost_rank"] = np.arange(1, len(out) + 1)
            info.update(applied=True, version=manifest["model_version"],
                        mode="CatBoost · " + (
                            "отложенный тест выполнен" if validation == "held_out_test" else "ЭКСПЕРИМЕНТ: без независимой проверки"),
                        model_sha256=manifest["model_sha256"], validation_status=validation,
                        schema_version=manifest["schema_version"])
            if manifest.get("catboost_version") != __version__:
                info[
                    "warning"] = "Версия библиотеки CatBoost отличается от версии при обучении; проверьте совместимость."
        except Exception as exc:
            out["catboost_score"] = np.nan
            out["catboost_rank"] = np.nan
            info.update(applied=False, version="", mode="CatBoost не применён: ошибка модели",
                        warning=f"{type(exc).__name__}: {exc}")
    elif model_path:
        info.update(mode="CatBoost не применён: файл модели не найден", warning=str(path))
    out["catboost_mode"] = info["mode"]
    out["catboost_version"] = info["version"]
    order = [c for c in
             ["source_inn", "candidate_inn", "catboost_rank", "catboost_score", "catboost_mode", "catboost_version"] if
             c in out]
    return out[order + [c for c in out.columns if c not in order]], info


def export_labeling(frame, target, metadata=None):
    """Экспорт только исходных признаков, без прогноза (снижает смещение разметки)."""
    import tempfile
    import os
    from openpyxl import load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    target = Path(target)
    columns = ALL_FEATURE_COLUMNS if all(c in frame.columns for c in ALL_FEATURE_COLUMNS) else FEATURE_COLUMNS
    data = frame.loc[:, ["source_inn", "candidate_inn"] + list(columns)].copy()
    meta = metadata or {}
    for c, value in {"analysis_id": meta.get("analysis_id", ""), "period_start": "", "period_end": "",
                     "data_version": "", LABEL: "", "expert_comment": "", "checked_at": ""}.items():
        data[c] = value
    instructions = [
        "1 — переход бизнеса подтверждён по принятому регламенту проверки; 0 — проверен и не подтвердился.",
        "Пустая метка — недостаточно сведений / не проверено. Она НЕ превращается в 0.",
        "Не копируйте оценку модели в метку и не размечайте по одному лишь числу совпадений.",
        "period_start и period_end: укажите общий сопоставимый период в формате ГГГГ-ММ-ДД для каждой размеченной строки.",
        "ИНН храните текстом. Запишите версию данных, комментарий и дату экспертной проверки.",
        "Объединяйте разметку разных исходных организаций и независимых случаев; нужны оба класса.",
        "Не меняйте рассчитанные признаки вручную. Незагруженное направление: пустые значения, не нули.",
        "При обучении исходные организации и кандидаты разделяются между обучением, валидацией и тестом без пересечения ИНН.",
        "Без модели баллы пустые. При наличии модели балл не является доказательством и не считается калиброванной вероятностью."]
    fd, temp = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
    os.close(fd)
    try:
        with pd.ExcelWriter(temp, engine="openpyxl") as writer:
            data.to_excel(writer, sheet_name="Для_разметки_CatBoost", index=False)
            pd.DataFrame({"Инструкция": instructions}).to_excel(writer, sheet_name="Инструкция", index=False)
        book = load_workbook(temp)
        for sheet in book:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        sheet = book["Для_разметки_CatBoost"]
        for row in sheet.iter_rows(min_row=2, max_row=max(2, sheet.max_row), min_col=1, max_col=2):
            for cell in row:
                cell.number_format = "@"
        dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
        dv.errorTitle, dv.error, dv.showErrorMessage = "Ошибка метки", "Только 0, 1 или пусто.", True
        sheet.add_data_validation(dv)
        column = sheet.cell(1, data.columns.get_loc(LABEL) + 1).column_letter
        dv.add(f"{column}2:{column}{max(2, sheet.max_row)}")
        book.save(temp)
        book.close()
        os.replace(temp, target)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
