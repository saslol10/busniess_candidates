import sys
from pathlib import Path

from PySide6.QtCore import Qt, QThread, Signal, QAbstractTableModel, QModelIndex, QObject, QEvent
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                               QHBoxLayout, QLabel, QPushButton, QLineEdit, QFileDialog, QMessageBox,
                               QTabWidget, QTableView, QPlainTextEdit, QSplitter, QDialog, QLayout, QScrollArea,
                               QSizePolicy, QHeaderView)

from manual_candidates_ui import CandidateSearchDialog
# action-column-only-v1
from catboost_action_column_only import attach_action_column
CandidateSearchDialog = attach_action_column(CandidateSearchDialog)

from core import load_book, valid_inn
from kkt_search import open_kkt_dialog
from staff import search_staff
from org_roles import search_roles
from trade_ui import TradeDialog


class CompactMessageBoxes(QObject):
    """Keep notifications narrow without changing search or table windows."""

    def eventFilter(self, obj, event):
        if isinstance(obj, QMessageBox) and event.type() == QEvent.Type.Show:
            screen = obj.screen()
            available = screen.availableGeometry() if screen else None
            width = 360 if len(obj.text()) < 100 else 420
            if available is not None:
                width = min(width, max(220, available.width() - 48))
            obj.setWindowState(Qt.WindowState.WindowNoState)
            obj.setMinimumSize(0, 0)
            obj.setMaximumSize(16777215, 16777215)
            layout = obj.layout()
            if layout:
                layout.setSizeConstraint(QLayout.SizeConstraint.SetNoConstraint)
            for label in obj.findChildren(QLabel):
                if label.objectName() in ('qt_msgbox_label', 'qt_msgbox_informativelabel'):
                    label.setWordWrap(True)
                    label.setMinimumWidth(0)
                    label.setMaximumWidth(max(160, width - 96))
            obj.setFixedWidth(width)
            if layout:
                layout.invalidate()
                layout.activate()
                height = layout.totalHeightForWidth(
                    width) if layout.hasHeightForWidth() else layout.totalSizeHint().height()
                obj.resize(width, max(120, height))
            if available is not None:
                parent = obj.parentWidget()
                center = parent.frameGeometry().center() if parent else available.center()
                x = max(available.left(), min(center.x() - obj.width() // 2, available.right() - obj.width() + 1))
                y = max(available.top(), min(center.y() - obj.height() // 2, available.bottom() - obj.height() + 1))
                obj.move(x, y)
        return super().eventFilter(obj, event)


class Loader(QThread):
    loaded = Signal(object)
    failed = Signal(str)

    def __init__(self, path):
        super().__init__()
        self.path = path

    def run(self):
        try:
            self.loaded.emit(load_book(self.path))
        except Exception as exc:
            self.failed.emit(str(exc))


class FrameModel(QAbstractTableModel):
    def __init__(self, frame, parent=None):
        super().__init__(parent)
        self.frame = frame

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.frame)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.frame.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if index.isValid() and role == Qt.ItemDataRole.DisplayRole:
            return str(self.frame.iat[index.row(), index.column()])
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return str(self.frame.columns[section]) if orientation == Qt.Orientation.Horizontal else str(section + 1)
        return None


class StaffDialog(QDialog):
    def __init__(self, result, source, parent=None):
        super().__init__(parent)
        self.result = result
        self.setWindowTitle("Общие сотрудники · ИНН " + source)
        self.resize(1150, 740)
        layout = QVBoxLayout(self)
        self.summary = QLabel(
            f"ИНН исходной: {source} · уникальных сотрудников: {result['source_count']} · кандидатов: {len(result['ranking'])} · пропущено строк: {result['skipped']}")
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        notice = QLabel("\n".join(result['warnings']))
        notice.setWordWrap(True)
        notice.setStyleSheet("background:#fff5dc;color:#614910;padding:12px;border-radius:6px")
        layout.addWidget(notice)
        actions = QHBoxLayout()
        export = QPushButton("Сохранить рейтинг и основания в Excel…")
        export.clicked.connect(self.export_result)
        actions.addWidget(export)
        close = QPushButton("Закрыть")
        close.clicked.connect(self.accept)
        actions.addStretch()
        actions.addWidget(close)
        layout.addLayout(actions)
        splitter = QSplitter(Qt.Orientation.Vertical)
        self.ranking = QTableView()
        self.ranking.setModel(FrameModel(result['ranking'], self.ranking))
        self.ranking.setEditTriggers(QTableView.EditTrigger.NoEditTriggers)
        self.ranking.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        self.ranking.setSelectionMode(QTableView.SelectionMode.SingleSelection)
        self.ranking.resizeColumnsToContents()
        splitter.addWidget(self.ranking)
        detail_widget = QWidget()
        detail_layout = QVBoxLayout(detail_widget)
        self.detail_label = QLabel("Выберите строку кандидата, чтобы увидеть основания.")
        self.detail_label.setWordWrap(True)
        detail_layout.addWidget(self.detail_label)
        self.details = QTableView()
        self.details.setEditTriggers(QTableView.EditTrigger.NoEditTriggers)
        self.details.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        detail_layout.addWidget(self.details)
        splitter.addWidget(detail_widget)
        splitter.setSizes([220, 350])
        layout.addWidget(splitter)
        self.ranking.selectionModel().currentRowChanged.connect(self.select_candidate)
        if not result['ranking'].empty:
            self.ranking.selectRow(0)
            self.select_candidate(self.ranking.model().index(0, 0), QModelIndex())
        else:
            self.detail_label.setText(
                "Кандидаты не найдены. Это не доказывает отсутствие перехода: нужны данные других организаций.")

    def select_candidate(self, current, previous):
        if not current.isValid():
            return
        inn = str(self.result['ranking'].iloc[current.row(), 0])
        self.detail_label.setText("Кандидат " + inn + " · основания из листа «Сотрудники»")
        self.details.setModel(FrameModel(self.result['details'][inn], self.details))
        self.details.resizeColumnsToContents()
        for c in range(self.details.model().columnCount()):
            self.details.setColumnWidth(c, min(360, max(150, self.details.columnWidth(c))))

    def export_result(self):
        path, _ = QFileDialog.getSaveFileName(self, "Сохранить результат", "staff_candidates.xlsx",
                                              "Книга Excel (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith('.xlsx'):
            path += '.xlsx'
        try:
            import pandas as pd
            import os, tempfile
            from openpyxl import load_workbook
            target = Path(path)
            fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=target.parent)
            os.close(fd)
            try:
                with pd.ExcelWriter(temp, engine='openpyxl') as writer:
                    self.result['ranking'].to_excel(writer, sheet_name='Кандидаты', index=False)
                    frames = []
                    for inn, frame in self.result['details'].items():
                        f = frame.copy()
                        f.insert(0, 'ИНН кандидата', inn)
                        frames.append(f)
                    if frames:
                        pd.concat(frames, ignore_index=True).to_excel(writer, sheet_name='Основания', index=False)
                    pd.DataFrame({'Методика': self.result['warnings']}).to_excel(writer, sheet_name='Методика',
                                                                                 index=False)
                book = load_workbook(temp)
                # Строки, начинающиеся со знака равенства, должны оставаться текстом.
                for sheet in book:
                    sheet.freeze_panes = 'A2'
                    sheet.auto_filter.ref = sheet.dimensions
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value, str):
                                cell.data_type = 's'
                book.save(temp)
                os.replace(temp, target)
            finally:
                if os.path.exists(temp): os.unlink(temp)
            QMessageBox.information(self, "Готово", "Рейтинг, основания и методика сохранены.")
        except Exception as exc:
            QMessageBox.warning(self, "Не удалось сохранить", str(exc))



class RoleDialog(QDialog):
    def __init__(self, result, source, parent=None):
        super().__init__(parent)
        self.result = result
        self.source = source
        self.setWindowTitle("Учредители, руководители и сотрудники · ИНН " + source)
        self.resize(1280, 760)
        layout = QVBoxLayout(self)
        self.summary = QLabel(result["summary"])
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        if result["warnings"]:
            notice = QLabel("\n".join(result["warnings"]))
            notice.setWordWrap(True)
            notice.setStyleSheet("background:#fff5dc;color:#614910;padding:12px;border-radius:6px")
            layout.addWidget(notice)
        actions = QHBoxLayout()
        export = QPushButton("Сохранить рейтинг и основания в Excel…")
        export.clicked.connect(self.export_result)
        actions.addWidget(export)
        actions.addStretch()
        close = QPushButton("Закрыть")
        close.clicked.connect(self.accept)
        actions.addWidget(close)
        layout.addLayout(actions)
        splitter = QSplitter(Qt.Orientation.Vertical)
        self.ranking = QTableView()
        self.ranking.setModel(FrameModel(result["ranking"], self.ranking))
        self.ranking.setEditTriggers(QTableView.EditTrigger.NoEditTriggers)
        self.ranking.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        self.ranking.setSelectionMode(QTableView.SelectionMode.SingleSelection)
        self.ranking.setAlternatingRowColors(True)
        self.ranking.horizontalHeader().setStretchLastSection(True)
        self.ranking.selectionModel().currentRowChanged.connect(self.select_candidate)
        splitter.addWidget(self.ranking)
        detail_widget = QWidget()
        detail_layout = QVBoxLayout(detail_widget)
        self.detail_label = QLabel("Выберите строку кандидата, чтобы увидеть основания.")
        self.detail_label.setWordWrap(True)
        detail_layout.addWidget(self.detail_label)
        self.details = QTableView()
        self.details.setEditTriggers(QTableView.EditTrigger.NoEditTriggers)
        self.details.setAlternatingRowColors(True)
        detail_layout.addWidget(self.details)
        splitter.addWidget(detail_widget)
        splitter.setSizes([430, 230])
        layout.addWidget(splitter, 1)
        if self.ranking.model().rowCount():
            self.ranking.selectRow(0)

    def select_candidate(self, current, previous):
        if not current.isValid():
            return
        inn = str(self.result["ranking"].iloc[current.row()]["ИНН кандидата"])
        name = str(self.result["ranking"].iloc[current.row()]["Наименование кандидата"])
        self.detail_label.setText("Кандидат " + inn + " · " + name + " · основания совпадений")
        self.details.setModel(FrameModel(self.result["details"].get(inn, self.result["empty_details"]), self.details))
        self.details.resizeColumnsToContents()
        for c in range(self.details.model().columnCount()):
            self.details.setColumnWidth(c, min(420, max(150, self.details.columnWidth(c))))

    def export_result(self):
        path, _ = QFileDialog.getSaveFileName(self, "Сохранить результат", "role_candidates.xlsx",
                                              "Книга Excel (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            import os, tempfile
            import pandas as pd
            from openpyxl import load_workbook
            target = Path(path)
            fd, temp = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
            os.close(fd)
            try:
                with pd.ExcelWriter(temp, engine="openpyxl") as writer:
                    self.result["ranking"].to_excel(writer, sheet_name="Кандидаты", index=False)
                    frames = []
                    for inn, frame in self.result["details"].items():
                        f = frame.copy()
                        f.insert(0, "ИНН кандидата", inn)
                        frames.append(f)
                    if frames:
                        pd.concat(frames, ignore_index=True).to_excel(writer, sheet_name="Основания", index=False)
                    pd.DataFrame({"Методика": self.result["warnings"] or ["Совпадения рассчитаны по ИНН физических лиц."]}).to_excel(
                        writer, sheet_name="Методика", index=False)
                book = load_workbook(temp)
                for sheet in book:
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value, str):
                                cell.data_type = "s"
                book.save(temp)
                os.replace(temp, target)
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)
            QMessageBox.information(self, "Готово", "Рейтинг и основания сохранены.")
        except Exception as exc:
            QMessageBox.warning(self, "Не удалось сохранить", str(exc))


class Window(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Перевод бизнеса")
        self.resize(1180, 780)
        self.frames = {}
        self.loader = None
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        title = QLabel("Перевод бизнеса")
        title.setStyleSheet("font-size:25px;font-weight:700;color:#173556;padding:8px 0")
        layout.addWidget(title)
        buttons = QHBoxLayout()
        self.open = QPushButton("Загрузить Excel…")
        self.open.clicked.connect(self.choose)
        buttons.addWidget(self.open)
        buttons.addStretch()
        layout.addLayout(buttons)
        self.file_label = QLabel("Файл не выбран")
        self.file_label.setWordWrap(True)
        layout.addWidget(self.file_label)
        filter_row = QHBoxLayout()
        self.inn = QLineEdit()
        self.inn.setPlaceholderText("10 цифр ИНН организации")
        self.inn.setMaxLength(10)
        self.inn.returnPressed.connect(self.apply_filter)
        self.filter_button = QPushButton("Показать записи организации")
        self.filter_button.clicked.connect(self.apply_filter)
        self.clear = QPushButton("Показать все")
        self.clear.clicked.connect(self.clear_filter)
        filter_row.addWidget(self.inn)
        filter_row.addWidget(self.filter_button)
        filter_row.addWidget(self.clear)
        layout.addLayout(filter_row)

        search_row = QHBoxLayout()

        self.staff_button = QPushButton("Поиск по сотрудникам")
        self.staff_button.setEnabled(False)
        self.staff_button.clicked.connect(self.find_staff)
        search_row.addWidget(self.staff_button)

        self.roles_button = QPushButton("Учредители и руководители…")
        self.roles_button.setToolTip("Сравнить сотрудников, учредителей и руководителей исходной организации с кандидатами")
        self.roles_button.setEnabled(False)
        self.roles_button.clicked.connect(self.find_roles)
        search_row.addWidget(self.roles_button)

        # separate-feature-windows-v1
        self.suppliers_button = QPushButton("Поставщики…")
        self.suppliers_button.setToolTip("Отдельное окно: общие поставщики и основания совпадений")
        self.suppliers_button.clicked.connect(lambda: self.open_trade("suppliers"))
        search_row.addWidget(self.suppliers_button)

        self.buyers_button = QPushButton("Покупатели…")
        self.buyers_button.setToolTip("Отдельное окно: общие покупатели и основания совпадений")
        self.buyers_button.clicked.connect(lambda: self.open_trade("buyers"))
        search_row.addWidget(self.buyers_button)

        self.kkt_button = QPushButton("Найти по ККТ")
        self.kkt_button.clicked.connect(lambda: open_kkt_dialog(self))
        search_row.addWidget(self.kkt_button)

        self.candidates_button = QPushButton("Общий поиск кандидатов…")
        self.candidates_button.setEnabled(False)
        self.candidates_button.setStyleSheet("background:#d7e7ff;font-weight:600;color:#0e294b")
        self.candidates_button.clicked.connect(self.open_candidates)
        search_row.addWidget(self.candidates_button)

        search_row.addStretch()
        layout.addLayout(search_row)
        self.status = QLabel("Загрузите заполненный файл.")
        self.status.setWordWrap(True)
        self.status.setStyleSheet("background:#eaf1f9;color:#173556;padding:12px;border-radius:6px")
        layout.addWidget(self.status)
        splitter = QSplitter(Qt.Orientation.Vertical)
        self.tabs = QTabWidget()
        self.log = QPlainTextEdit()
        self.log.setReadOnly(True)
        splitter.addWidget(self.tabs)
        splitter.addWidget(self.log)
        splitter.setSizes([440, 160])
        layout.addWidget(splitter, 1)
        self.setStyleSheet("""
            QWidget{font-family:Segoe UI;font-size:13px;color:#182b41;background:#f4f7fb}
            QPushButton{padding:9px 12px;background:#e6eefb;border:1px solid #acbdd6;border-radius:5px;color:#173556}
            QPushButton:hover{background:#d5e4fc} QPushButton:disabled{color:#657387;background:#e8ecf1}
            QLineEdit,QPlainTextEdit,QTableView{background:#ffffff;color:#182b41;border:1px solid #b9c8da;selection-background-color:#d7e7ff;selection-color:#102d50}
            QLineEdit{padding:8px} QTableView{gridline-color:#e0e7f0;alternate-background-color:#eef3fa}
            QHeaderView::section{background:#e7eef8;color:#173556;padding:7px;border:1px solid #cad6e5}
            QTabBar::tab{background:#e5edf8;color:#173556;padding:8px 12px}
            QTabBar::tab:selected{background:#ffffff;border-bottom:2px solid #345fe2}
            QToolTip{background:#ffffff;color:#182b41;border:1px solid #b9c8da}
        """)

    def open_candidates(self):
        if not self.frames:
            QMessageBox.information(self, "Нет данных", "Сначала загрузите Excel.")
            return
        inn = self.inn.text().strip()
        if not valid_inn(inn, (10,)):
            QMessageBox.warning(self, "ИНН", "Введите корректный 10-значный ИНН организации.")
            return
        CandidateSearchDialog(self.frames, inn, self).exec()

    def choose(self):
        path, _ = QFileDialog.getOpenFileName(self, "Выберите файл Excel", "", "Книги Excel (*.xlsx)")
        if path:
            self.start(path)

    def start(self, path):
        if self.loader and self.loader.isRunning():
            return
        self.frames = {}
        self.staff_button.setEnabled(False)
        self.roles_button.setEnabled(False)
        self.candidates_button.setEnabled(False)
        self.tabs.clear()
        self.log.clear()
        self.open.setEnabled(False)
        self.file_label.setText("Файл: " + Path(path).name)
        self.status.setText("Чтение файла и проверка данных…")
        self.loader = Loader(path)
        self.loader.loaded.connect(self.loaded)
        self.loader.failed.connect(self.failed)
        self.loader.finished.connect(self.finished)
        self.loader.start()

    def finished(self):
        self.open.setEnabled(True)
        self.staff_button.setEnabled("Сотрудники" in self.frames)
        self.roles_button.setEnabled(bool(self.frames))
        self.candidates_button.setEnabled(bool(self.frames))

    def failed(self, message):
        self.status.setText("Файл не загружен.")
        self.log.setPlainText(message)
        QMessageBox.warning(self, "Ошибка загрузки", message)

    def loaded(self, result):
        self.frames = result["frames"]
        count = sum(map(len, self.frames.values()))
        orgs = set()
        for frame in self.frames.values():
            orgs.update(frame["ИНН организации"].tolist())
        self.status.setText(
            f"Листов: {len(self.frames)} · строк: {count} · ИНН организаций: {len(orgs)} · ошибок: {result['error_count']}. " + (
                "Исправьте ошибки в Excel и загрузите файл повторно." if result[
                    'error_count'] else "Структурная проверка пройдена."))
        parts = ["ОШИБКИ", *(result["errors"] or ["Не обнаружены."]), "", "ПРЕДУПРЕЖДЕНИЯ И ОГРАНИЧЕНИЯ",
                 *result["warnings"]]
        self.log.setPlainText("\n".join(parts))
        self.inn.clear()
        self.render()

    def open_trade(self, mode=None):
        if not self.frames:
            QMessageBox.information(self, "Нет данных", "Сначала загрузите Excel.")
            return
        dialog = TradeDialog(self.frames, self.inn.text().strip(), self)
        if mode is not None:
            index = dialog.mode.findData(mode)
            if index < 0:
                QMessageBox.warning(self, "Несовместимая версия", "В trade_ui.py не найден нужный режим поиска.")
                return
            dialog.mode.setCurrentIndex(index)
            dialog.mode.setEnabled(False)
            dialog.setWindowTitle(
                "Общие поставщики · отдельный признак" if mode == "suppliers"
                else "Общие покупатели · отдельный признак"
            )
        dialog.exec()

    def render(self, inn=""):
        self.tabs.clear()
        for name, frame in self.frames.items():
            viewframe = frame.loc[frame["ИНН организации"] == inn] if inn else frame
            table = QTableView()
            table.setModel(FrameModel(viewframe.reset_index(drop=True), table))
            table.setAlternatingRowColors(True)
            table.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
            for c in range(len(viewframe.columns)):
                table.setColumnWidth(c, 170 if c else 145)
            self.tabs.addTab(table, f"{name} ({len(viewframe)})")

    def apply_filter(self):
        if not self.frames:
            QMessageBox.information(self, "Нет данных", "Сначала загрузите Excel.")
            return
        inn = self.inn.text().strip()
        if not valid_inn(inn, (10,)):
            QMessageBox.warning(self, "Проверьте ИНН",
                                "Нужен 10-значный ИНН организации с корректной контрольной цифрой.")
            return
        self.render(inn)
        count = sum(int((f["ИНН организации"] == inn).sum()) for f in self.frames.values())
        self.log.appendPlainText(f"\nФильтр {inn}: найдено строк — {count}.")

    def clear_filter(self):
        self.inn.clear()
        self.render()

    def find_staff(self):
        try:
            result = search_staff(self.frames.get("Сотрудники"), self.inn.text().strip())
            StaffDialog(result, self.inn.text().strip(), self).exec()
        except Exception as exc:
            QMessageBox.warning(self, "Поиск по сотрудникам", str(exc))

    def find_roles(self):
        try:
            source = self.inn.text().strip()
            result = search_roles(self.frames, source)
            RoleDialog(result, source, self).exec()
        except Exception as exc:
            QMessageBox.warning(self, "Учредители и руководители", str(exc))

    def closeEvent(self, event):
        if self.loader and self.loader.isRunning():
            QMessageBox.information(self, "Загрузка", "Дождитесь завершения чтения файла перед закрытием.")
            event.ignore()
        else:
            event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    message_box_filter = CompactMessageBoxes(app)
    app.installEventFilter(message_box_filter)
    window = Window()
    window.show()
    sys.exit(app.exec())
